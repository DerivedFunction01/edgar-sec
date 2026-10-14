#!/usr/bin/env python3
"""
Zero-Pandas Excel Audit Workbook Builder (`build_audit_workbook.py`).
Builds a clean, formatted multi-tab `.xlsx` workbook from the v3.0 evaluation outputs
using `openpyxl` with CLI arguments (`--data-dir`, `--output-dir`, `--output-file`).
"""

import argparse
import csv
import json
import pathlib
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def auto_fit_and_style_sheet(ws: Any, header_fill_hex: str = "1F4E78") -> None:
    header_fill = PatternFill(start_color=header_fill_hex, end_color=header_fill_hex, fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 26

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 20
        for col_idx in range(1, ws.max_column + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = body_font
            c.border = thin_border
            if isinstance(c.value, (int, float)):
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)


def add_csv_sheet(wb: openpyxl.Workbook, title: str, csv_path: pathlib.Path, header_hex: str = "1F4E78") -> None:
    ws = wb.create_sheet(title=title[:31])
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        for r_idx, row in enumerate(reader, start=1):
            parsed_row = []
            for val in row:
                if r_idx > 1:
                    try:
                        if "." in val:
                            parsed_row.append(float(val))
                        else:
                            parsed_row.append(int(val))
                        continue
                    except ValueError:
                        pass
                parsed_row.append(val)
            ws.append(parsed_row)
    auto_fit_and_style_sheet(ws, header_hex)


def build_workbook(data_dir: pathlib.Path, output_dir: pathlib.Path, out_file: pathlib.Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    with open(output_dir / "section9_evaluation_summary.json", "r", encoding="utf-8") as f:
        summary = json.load(f)

    # Sheet 1: Section 9 Executive Metrics
    ws_s9 = wb.create_sheet(title="Section9_Metrics_Summary")
    ws_s9.append([
        "Cohort Split", "Total Facts", "Fact Precision (%)", "Fact Coverage (%)",
        "Placement Accuracy (%)", "Unknown Abstention Rate (%)",
        "Total Goals", "Goals Solved", "Goal Precision (%)", "Goal Coverage (%)",
        "Multi-Hop Derived Goals Solved", "Multi-Hop Derived Total", "Multi-Hop Coverage (%)",
        "Direct Lookup Goals Solved", "Direct Lookup Total", "Direct Lookup Coverage (%)"
    ])
    for key in ["overall", "seed_cohort", "holdout_zero_shot_cohort"]:
        m = summary["section9_metrics"][key]
        ws_s9.append([
            m["cohort"], m["total_facts"], m["fact_precision_pct"], m["fact_coverage_recall_pct"],
            m["placement_accuracy_pct"], m["unknown_abstention_rate_pct"],
            m["total_goals"], m["goals_solved_correctly"], m["goal_precision_pct"], m["goal_coverage_recall_pct"],
            m["derived_multi_hop_goals"]["correct"], m["derived_multi_hop_goals"]["total"], m["derived_multi_hop_goals"]["coverage_pct"],
            m["direct_lookup_goals"]["correct"], m["direct_lookup_goals"]["total"], m["direct_lookup_goals"]["coverage_pct"],
        ])
    auto_fit_and_style_sheet(ws_s9, "1F4E78")

    # Sheet 2: 4-Way Ablation Study
    add_csv_sheet(wb, "Ablation_Study_4Way", output_dir / "ablation_study_section9.csv", "2F5597")
    # Sheet 3: Accounting Identity Oracle Checks
    add_csv_sheet(wb, "Identity_Oracle_Checks", output_dir / "accounting_identity_oracle_checks.csv", "375623")
    # Sheet 4: Parity & Goal Verdicts
    add_csv_sheet(wb, "Goal_Verdicts_And_Parity", output_dir / "parity_and_oracle_verdicts.csv", "1F4E78")
    # Sheet 5: Unknown Resolution Queue
    add_csv_sheet(wb, "Unknown_Resolution_Queue", output_dir / "unknown_resolution_queue.csv", "833C0C")
    # Sheet 6–9: Reusable KB Tables
    add_csv_sheet(wb, "KB_Prime_Heads", data_dir / "kb_prime_heads.csv", "404040")
    add_csv_sheet(wb, "KB_SubAtomic_Modifiers", data_dir / "kb_subatomic_modifiers.csv", "404040")
    add_csv_sheet(wb, "KB_Typed_Relations", data_dir / "kb_typed_relations.csv", "404040")
    add_csv_sheet(wb, "KB_Geography_DAG", data_dir / "kb_geography_dag.csv", "404040")

    wb.save(out_file)
    print(f"[EXCEL WORKBOOK] Built {out_file.name} with {len(wb.sheetnames)} audit tabs (Zero Pandas).")


def main() -> None:
    base_dir = pathlib.Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(description="Build Clean-Room v3.0 Multi-Tab Excel Audit Workbook")
    parser.add_argument("--data-dir", type=pathlib.Path, default=base_dir / "data", help="Input data directory")
    parser.add_argument("--output-dir", type=pathlib.Path, default=base_dir / "outputs", help="Evaluation outputs directory")
    parser.add_argument(
        "--output-file",
        type=pathlib.Path,
        default=base_dir / "outputs" / "xbrl_clean_engine_v3_audit.xlsx",
        help="Path to output .xlsx workbook",
    )
    args = parser.parse_args()
    build_workbook(args.data_dir, args.output_dir, args.output_file)


if __name__ == "__main__":
    main()
